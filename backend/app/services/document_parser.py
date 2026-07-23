from html.parser import HTMLParser
from pathlib import Path
import re

from docx import Document as DocxDocument
from markdown import markdown
from pypdf import PdfReader

from app.services.text_splitter import TextBlock


class DocumentParseError(ValueError):
    """文档内容无法被解析。"""


class _HtmlTextExtractor(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        text = data.strip()
        if text:
            self.parts.append(text)

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"p", "br", "li", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote"}:
            self.parts.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"p", "li", "h1", "h2", "h3", "h4", "h5", "h6", "pre", "blockquote"}:
            self.parts.append("\n")

    def get_text(self) -> str:
        lines: list[str] = []
        current: list[str] = []
        for part in self.parts:
            if part == "\n":
                if current:
                    lines.append(" ".join(current).strip())
                    current = []
                continue
            current.append(part)
        if current:
            lines.append(" ".join(current).strip())
        return "\n".join(line for line in lines if line).strip()


class DocumentParser:
    supported_types = {"pdf", "docx", "md", "markdown", "txt"}

    def parse(self, file_path: str, file_type: str) -> list[TextBlock]:
        path = Path(file_path)
        normalized_type = file_type.lower().lstrip(".")

        try:
            if normalized_type == "pdf":
                blocks = self._parse_pdf(path)
            elif normalized_type == "docx":
                blocks = self._parse_docx(path)
            elif normalized_type in {"md", "markdown"}:
                blocks = self._parse_markdown(path)
            elif normalized_type == "txt":
                blocks = [TextBlock(text=self._read_text(path))]
            else:
                raise DocumentParseError(f"不支持的文档格式：{file_type}。")
        except DocumentParseError:
            raise
        except FileNotFoundError as exc:
            raise DocumentParseError("文档解析失败：找不到已上传的文件，请重新上传。") from exc
        except Exception as exc:
            raise DocumentParseError(
                "文档解析失败：文件可能已损坏、被加密或内容格式不正确。"
            ) from exc

        if not any(block.text.strip() for block in blocks):
            raise DocumentParseError(
                "文档解析失败：未提取到可用文字；扫描版 PDF 请先进行 OCR 识别。"
            )
        return blocks

    def _parse_pdf(self, path: Path) -> list[TextBlock]:
        reader = PdfReader(str(path))
        blocks: list[TextBlock] = []

        for page_index, page in enumerate(reader.pages, start=1):
            text = page.extract_text() or ""
            if text.strip():
                blocks.append(TextBlock(text=text, page_no=page_index))

        return blocks

    def _parse_docx(self, path: Path) -> list[TextBlock]:
        document = DocxDocument(str(path))
        blocks: list[TextBlock] = []
        heading_stack: list[str] = []
        current_lines: list[str] = []
        current_title: str | None = None

        def flush() -> None:
            nonlocal current_lines
            text = "\n".join(current_lines).strip()
            if text:
                blocks.append(TextBlock(text=text, section_title=current_title))
            current_lines = []

        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            style_name = str(getattr(paragraph.style, "name", "") or "")
            match = re.search(r"(?:Heading|标题)\s*(\d+)", style_name, re.IGNORECASE)
            if match:
                flush()
                level = max(1, int(match.group(1)))
                heading_stack[level - 1 :] = [value]
                current_title = " > ".join(heading_stack)
            else:
                current_lines.append(value)
        for table in document.tables:
            for row in table.rows:
                values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if values:
                    current_lines.append(" | ".join(values))
        flush()
        return blocks

    def _parse_markdown(self, path: Path) -> list[TextBlock]:
        raw = self._read_text(path)
        sections: list[TextBlock] = []
        heading_stack: list[str] = []
        current_lines: list[str] = []
        current_title: str | None = None
        fence: str | None = None

        def extract(value: str) -> str:
            html = markdown(value)
            extractor = _HtmlTextExtractor()
            extractor.feed(html)
            return extractor.get_text() or value.strip()

        def flush() -> None:
            nonlocal current_lines
            value = "\n".join(current_lines).strip()
            if value:
                sections.append(
                    TextBlock(text=extract(value), section_title=current_title)
                )
            current_lines = []

        for line in raw.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
            stripped = line.lstrip()
            marker = stripped[:3]
            if marker in {"```", "~~~"}:
                fence = None if fence == marker else marker
                current_lines.append(line)
                continue
            match = None if fence else re.match(r"^\s{0,3}(#{1,6})\s+(.+?)\s*#*\s*$", line)
            if match:
                flush()
                level = len(match.group(1))
                title = match.group(2).strip()
                heading_stack[level - 1 :] = [title]
                current_title = " > ".join(heading_stack)
                continue
            current_lines.append(line)
        flush()
        if sections:
            return sections
        html = markdown(raw)
        extractor = _HtmlTextExtractor()
        extractor.feed(html)
        extracted = extractor.get_text()
        return [TextBlock(text=extracted or raw)]

    def _read_text(self, path: Path) -> str:
        for encoding in ("utf-8", "utf-8-sig", "gb18030"):
            try:
                return path.read_text(encoding=encoding)
            except UnicodeDecodeError:
                continue
        return path.read_text(encoding="utf-8", errors="ignore")


document_parser = DocumentParser()
